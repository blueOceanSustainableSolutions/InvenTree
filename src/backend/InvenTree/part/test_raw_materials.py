"""Tests for the raw materials required by a BOM or a build order."""

from decimal import Decimal
from io import BytesIO

from django.urls import reverse

import openpyxl

from build.models import Build, BuildLine
from company.models import Company, SupplierPart
from InvenTree.unit_test import InvenTreeAPITestCase
from part.models import BomItem, Part
from part.raw_materials import (
    NOTE_CONSUMABLE,
    NOTE_NO_BOM,
    explode_build,
    explode_part,
    summarize,
)


class RawMaterialsTests(InvenTreeAPITestCase):
    """A unit made of boards, with setup, attrition and rounding on the BOM."""

    roles = ['part.view', 'build.view']

    @classmethod
    def setUpTestData(cls):
        """Unit = 2 boards (+1 setup), 3 screws, glue, a frame and rope."""
        super().setUpTestData()

        def part(name, **kwargs):
            return Part.objects.create(name=name, component=True, **kwargs)

        part('Screw', purchaseable=True)
        part('PCB')
        part('Glue')
        part('Rope', units='m')
        part('Board', assembly=True)
        part('Frame', assembly=True)
        part('Unit', assembly=True)

        # Part trees are renumbered as parts are added, so load them afterwards
        (cls.screw, cls.pcb, cls.glue, cls.rope, cls.board, cls.frame, cls.unit) = (
            Part.objects.get(name=name)
            for name in ('Screw', 'PCB', 'Glue', 'Rope', 'Board', 'Frame', 'Unit')
        )

        supplier = Company.objects.create(name='Acme', is_supplier=True)
        SupplierPart.objects.create(
            part=cls.screw, supplier=supplier, SKU='SCR-1', primary=True
        )

        def bom(parent, sub, quantity, **kwargs):
            return BomItem.objects.create(
                part=parent, sub_part=sub, quantity=quantity, **kwargs
            )

        bom(cls.board, cls.screw, 2, setup_quantity=4, reference='S1')
        bom(cls.board, cls.pcb, 1, attrition=10)
        bom(cls.unit, cls.board, 2, setup_quantity=1)
        bom(cls.unit, cls.screw, 3)
        bom(cls.unit, cls.glue, Decimal('0.5'), consumable=True)
        bom(cls.unit, cls.frame, 1)
        bom(cls.unit, cls.rope, Decimal('1.5'), rounding_multiple=2)

    def totals(self, rows):
        """Map part name to total quantity and notes."""
        return {
            entry['part'].name: (entry['quantity'], entry['notes'])
            for entry in summarize(rows)
        }

    def test_part_breakdown(self):
        """Sub-assemblies are replaced by their BOM, with setup once per level."""
        totals = self.totals(explode_part(self.unit, 3))
        self.assertEqual(set(totals), {'Screw', 'PCB', 'Glue', 'Frame', 'Rope'})
        # 2 * 3 + 1 = 7 boards: 2 * 7 + 4 screws, plus 3 * 3 directly
        self.assertEqual(totals['Screw'][0], Decimal(27))
        # 7 PCBs with 10 % attrition
        self.assertEqual(totals['PCB'][0], Decimal('7.7'))
        # 4.5 m rounded up to a multiple of 2
        self.assertEqual(totals['Rope'][0], Decimal(6))
        self.assertEqual(totals['Glue'], (Decimal('1.5'), {str(NOTE_CONSUMABLE)}))
        self.assertEqual(totals['Frame'], (Decimal(3), {str(NOTE_NO_BOM)}))

    def test_detail_path(self):
        """Each row records which assemblies it came through."""
        rows = explode_part(self.unit, 3)
        screws = [row for row in rows if row['part'] == self.screw]
        self.assertEqual(
            sorted((row['path'], row['quantity'], row['reference']) for row in screws),
            [([], Decimal(3 * 3), ''), (['Board x7'], Decimal(18), 'S1')],
        )

    def test_build_uses_its_own_lines(self):
        """A build order's lines, including build-only parts, are broken down."""
        build = Build.objects.create(part=self.unit, reference='BO-8101', quantity=3)
        BuildLine.objects.create(build=build, custom_part=self.pcb, quantity=2)
        # Lines can differ from the BOM: the build's quantity wins
        line = BuildLine.objects.get(build=build, bom_item__sub_part=self.glue)
        line.quantity = 4
        line.save()

        totals = self.totals(explode_build(build))
        self.assertEqual(totals['Screw'][0], Decimal(27))
        self.assertEqual(totals['PCB'][0], Decimal('9.7'))
        self.assertEqual(totals['Glue'][0], Decimal(4))

    def test_part_download(self):
        """The API returns a workbook with a summary and a detail sheet."""
        url = reverse('api-part-raw-materials', kwargs={'pk': self.unit.pk})
        response = self.client.get(url, {'quantity': 3})
        self.assertEqual(response.status_code, 200)
        self.assertIn('raw_materials_Unit_x3.xlsx', response['Content-Disposition'])

        workbook = openpyxl.load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Detail'])
        summary = {
            row[0]: row
            for row in workbook['Summary'].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(summary['Screw'][3], 27)
        self.assertEqual(summary['Screw'][5:7], ('Acme', 'SCR-1'))
        self.assertEqual(summary['Rope'][4], 'm')

    def test_part_download_rejects_bad_quantity(self):
        """The quantity must be positive."""
        url = reverse('api-part-raw-materials', kwargs={'pk': self.unit.pk})
        self.assertEqual(self.client.get(url, {'quantity': 0}).status_code, 400)

    def test_build_download(self):
        """Build orders can be exported too."""
        build = Build.objects.create(part=self.unit, reference='BO-8102', quantity=1)
        url = reverse('api-build-raw-materials', kwargs={'pk': build.pk})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('raw_materials_BO-8102.xlsx', response['Content-Disposition'])

    def test_requires_roles(self):
        """Users without the part or build role cannot download."""
        build = Build.objects.create(part=self.unit, reference='BO-8103', quantity=1)
        self.clearRoles()
        for url in (
            reverse('api-part-raw-materials', kwargs={'pk': self.unit.pk}),
            reverse('api-build-raw-materials', kwargs={'pk': build.pk}),
        ):
            self.assertEqual(self.client.get(url).status_code, 403)
