"""Raw materials of a built stock item, from the stock that was actually consumed.

Starting from a stock item, everything its build (or quick build) consumed is
listed. Any consumed item which was itself built is replaced by what *its* build
consumed, scaled by the share of that build's output which was used, until only
stock which was not built (i.e. bought) remains.

Quantities are proportional: if a build produced 10 boards and 3 were consumed,
3/10 of every untracked item that build consumed is attributed to them. Parts
installed into a serialized output belong to that output only.
"""

from collections import OrderedDict
from decimal import Decimal
from io import BytesIO

from django.utils.translation import gettext_lazy as _

from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import extend_schema
from rest_framework.generics import GenericAPIView

from InvenTree.helpers import DownloadFile, normalize
from stock.models import StockItem

# Guards against corrupt data (e.g. a build which consumed its own output)
MAX_DEPTH = 25

NOTE_NO_BUILD = _('Assembly without build record')
NOTE_NO_CONSUMED = _('Build has no consumed stock recorded')
NOTE_CYCLE = _('Circular build history')

RELATED = ('part', 'supplier_part', 'supplier_part__supplier', 'build', 'quick_build')


def _output_quantity(item: StockItem) -> Decimal:
    """Total quantity produced by the build which produced this item."""
    if item.build_id:
        quantity = item.build.completed or item.build.quantity
    else:
        quantity = item.quick_build.quantity
    return Decimal(quantity) if quantity else Decimal(1)


def _children(item: StockItem):
    """Yield (consumed item, factor) for the stock which went into this item.

    The factor is the share of the consumed item attributable to this item.
    """
    for child in item.installed_parts.select_related(*RELATED).order_by('pk'):
        yield child, Decimal(1)

    share = Decimal(item.quantity) / _output_quantity(item)

    if item.build_id:
        shared = StockItem.objects.filter(
            consumed_by=item.build_id, belongs_to__isnull=True
        )
    else:
        shared = StockItem.objects.filter(consumed_by_quick_build=item.quick_build_id)

    for child in shared.select_related(*RELATED).order_by('pk'):
        yield child, share


def _source(item: StockItem):
    """The build or quick build which produced this item, if any."""
    return item.build or item.quick_build


def explode(item: StockItem) -> list[dict]:
    """Return the raw material rows for the given stock item.

    Each row is one stock item which was not built, with the quantity used.
    Built items whose consumption cannot be traced are kept as rows with a note.
    """
    rows = []

    def walk(node, factor, path, seen, depth):
        source = _source(node)
        note = None

        if source is None:
            if node.part.assembly and depth > 0:
                note = NOTE_NO_BUILD
        elif (type(source), source.pk) in seen or depth >= MAX_DEPTH:
            note = NOTE_CYCLE
        else:
            children = list(_children(node))
            if children:
                key = (type(source), source.pk)
                for child, share in children:
                    walk(child, factor * share, [*path, node], seen | {key}, depth + 1)
                return
            note = NOTE_NO_CONSUMED

        if depth == 0 and note is None:
            # The starting item was not built: nothing to break down
            return

        rows.append({
            'item': node,
            'quantity': Decimal(node.quantity) * factor,
            'path': path,
            'note': str(note) if note else '',
        })

    root = StockItem.objects.select_related(*RELATED).get(pk=item.pk)
    walk(root, Decimal(1), [], frozenset(), 0)
    return rows


def _label(item: StockItem) -> str:
    """Short description of a stock item for the 'via' column."""
    if item.serial:
        return f'{item.part.full_name} #{item.serial}'
    if item.batch:
        return f'{item.part.full_name} ({item.batch})'
    return item.part.full_name


def summarize(rows: list[dict]) -> list[dict]:
    """Total the raw material rows per part."""
    totals = OrderedDict()

    for row in rows:
        part = row['item'].part
        entry = totals.setdefault(
            part.pk, {'part': part, 'quantity': Decimal(0), 'notes': set()}
        )
        entry['quantity'] += row['quantity']
        if row['note']:
            entry['notes'].add(row['note'])

    return sorted(totals.values(), key=lambda entry: entry['part'].full_name.lower())


XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def write_sheet(sheet, title, headers, data):
    """Fill a worksheet with a bold, frozen header row and sized columns."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheet.title = str(title)
    sheet.append([str(header) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for values in data:
        sheet.append(values)
    sheet.freeze_panes = 'A2'
    for idx, _header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        width = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value), default=8
        )
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 60)


def units(part) -> str:
    """The part's units, or blank."""
    return str(part.units) if part.units else ''


def save_workbook(workbook) -> bytes:
    """Return the workbook as file data."""
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_workbook(item: StockItem, rows: list[dict]) -> bytes:
    """Create the Excel workbook with a summary and a detail sheet."""
    import openpyxl

    workbook = openpyxl.Workbook()

    write_sheet(
        workbook.active,
        _('Summary'),
        [_('Part'), _('IPN'), _('Description'), _('Quantity'), _('Units'), _('Note')],
        [
            [
                entry['part'].full_name,
                entry['part'].IPN or '',
                entry['part'].description or '',
                float(normalize(entry['quantity'])),
                units(entry['part']),
                '; '.join(sorted(entry['notes'])),
            ]
            for entry in summarize(rows)
        ],
    )

    def supplier(stock):
        if stock.supplier_part:
            return stock.supplier_part.supplier.name, stock.supplier_part.SKU
        return '', ''

    write_sheet(
        workbook.create_sheet(),
        _('Detail'),
        [
            _('Part'),
            _('IPN'),
            _('Stock Item'),
            _('Serial'),
            _('Batch'),
            _('Supplier'),
            _('SKU'),
            _('Quantity'),
            _('Units'),
            _('Via'),
            _('Note'),
        ],
        [
            [
                row['item'].part.full_name,
                row['item'].part.IPN or '',
                row['item'].pk,
                row['item'].serial or '',
                row['item'].batch or '',
                *supplier(row['item']),
                float(normalize(row['quantity'])),
                units(row['item'].part),
                ' > '.join(_label(node) for node in row['path']),
                row['note'],
            ]
            for row in rows
        ],
    )

    return save_workbook(workbook)


class StockItemRawMaterials(GenericAPIView):
    """Download the raw materials consumed to build a stock item, as Excel."""

    queryset = StockItem.objects.all()
    # There is no serializer to infer the model from, so name the role
    role_required = 'stock'

    @extend_schema(responses={200: OpenApiTypes.BINARY})
    def get(self, request, *args, **kwargs):
        """Return the raw materials workbook."""
        item = self.get_object()
        data = build_workbook(item, explode(item))
        name = item.serial or item.batch or str(item.pk)
        filename = f'raw_materials_{item.part.name}_{name}.xlsx'.replace(' ', '_')
        return DownloadFile(data, filename, content_type=XLSX_CONTENT_TYPE)
