"""Tests for the raw materials breakdown of built stock items."""

from decimal import Decimal
from io import BytesIO

from django.urls import reverse

import openpyxl

from build.models import Build, QuickBuild
from build.status_codes import BuildStatus
from company.models import Company, SupplierPart
from InvenTree.unit_test import InvenTreeAPITestCase
from part.models import Part
from stock.models import StockItem
from stock.raw_materials import NOTE_CYCLE, NOTE_NO_BUILD, explode, summarize


class RawMaterialsTests(InvenTreeAPITestCase):
    """Two levels of builds, a quick build and untraceable stock."""

    roles = ['stock.view']

    @classmethod
    def setUpTestData(cls):
        """Build 10 boards, then 2 units which use 3 of them."""
        super().setUpTestData()

        def part(name, **kwargs):
            return Part.objects.create(name=name, component=True, **kwargs)

        cls.screw = part('Screw', purchaseable=True)
        cls.pcb = part('PCB')
        cls.motor = part('Motor')
        cls.rope = part('Rope', units='m')
        cls.board = part('Board', assembly=True)
        cls.cable = part('Cable', assembly=True)
        cls.bracket = part('Bracket', assembly=True)
        cls.product = part('Unit', assembly=True, trackable=True)

        supplier = Company.objects.create(name='Acme', is_supplier=True)
        cls.screw_sp = SupplierPart.objects.create(
            part=cls.screw, supplier=supplier, SKU='SCR-1'
        )

        def build(assembly, reference, completed):
            return Build.objects.create(
                part=assembly,
                reference=reference,
                quantity=completed,
                completed=completed,
                status=BuildStatus.COMPLETE,
            )

        # Level 1: 10 boards from 20 screws and 10 PCBs
        cls.board_build = build(cls.board, 'BO-8001', 10)
        StockItem.objects.create(
            part=cls.screw,
            quantity=20,
            consumed_by=cls.board_build,
            supplier_part=cls.screw_sp,
        )
        StockItem.objects.create(part=cls.pcb, quantity=10, consumed_by=cls.board_build)

        # Quick build: 2 cables from 4 m of rope
        cls.quick_build = QuickBuild.objects.create(
            reference='QB-8001', part=cls.cable, quantity=2
        )
        StockItem.objects.create(
            part=cls.rope, quantity=4, consumed_by_quick_build=cls.quick_build
        )

        # Level 2: 2 units sharing 3 boards, 4 screws, 1 cable and 2 brackets
        cls.unit_build = build(cls.product, 'BO-8002', 2)
        StockItem.objects.create(
            part=cls.board, quantity=7, build=cls.board_build, batch='B-1'
        )
        StockItem.objects.create(
            part=cls.board,
            quantity=3,
            build=cls.board_build,
            batch='B-1',
            consumed_by=cls.unit_build,
        )
        StockItem.objects.create(part=cls.screw, quantity=4, consumed_by=cls.unit_build)
        StockItem.objects.create(
            part=cls.cable,
            quantity=1,
            quick_build=cls.quick_build,
            consumed_by=cls.unit_build,
        )
        StockItem.objects.create(
            part=cls.bracket, quantity=2, consumed_by=cls.unit_build
        )

        cls.unit = StockItem.objects.create(
            part=cls.product, quantity=1, serial='U-1', build=cls.unit_build
        )
        StockItem.objects.create(
            part=cls.product, quantity=1, serial='U-2', build=cls.unit_build
        )
        # Installed into U-1 only
        StockItem.objects.create(
            part=cls.motor, quantity=1, belongs_to=cls.unit, consumed_by=cls.unit_build
        )

    def totals(self, item):
        """Map part name to total quantity and notes."""
        return {
            entry['part'].name: (entry['quantity'], entry['notes'])
            for entry in summarize(explode(item))
        }

    def test_breakdown_replaces_built_items(self):
        """Built items are replaced by their consumed stock, proportionally."""
        totals = self.totals(self.unit)
        self.assertEqual(set(totals), {'Screw', 'PCB', 'Motor', 'Rope', 'Bracket'})
        # Half of 3 boards = 1.5 boards: 3 screws + half of the 4 shared screws
        self.assertEqual(totals['Screw'][0], Decimal(5))
        self.assertEqual(totals['PCB'][0], Decimal('1.5'))
        # Installed only into this unit
        self.assertEqual(totals['Motor'][0], Decimal(1))
        # Half of 1 cable, which is half of 4 m of rope
        self.assertEqual(totals['Rope'][0], Decimal(1))
        # An assembly without a build is kept, and flagged
        self.assertEqual(totals['Bracket'][0], Decimal(1))
        self.assertEqual(totals['Bracket'][1], {str(NOTE_NO_BUILD)})

    def test_detail_path_and_supplier(self):
        """Each row records which assemblies it came through."""
        rows = [row for row in explode(self.unit) if row['item'].supplier_part]
        self.assertEqual(len(rows), 1)
        self.assertEqual(
            [node.part.name for node in rows[0]['path']], ['Unit', 'Board']
        )
        self.assertEqual(rows[0]['quantity'], Decimal(3))

    def test_item_which_was_not_built(self):
        """Bought stock has nothing to break down."""
        item = StockItem.objects.create(part=self.screw, quantity=5)
        self.assertEqual(explode(item), [])

    def test_circular_history_stops(self):
        """A build which consumed its own output does not recurse forever."""
        build = Build.objects.create(
            part=self.board,
            reference='BO-8003',
            quantity=1,
            completed=1,
            status=BuildStatus.COMPLETE,
        )
        item = StockItem.objects.create(
            part=self.board, quantity=1, build=build, consumed_by=build
        )
        rows = explode(item)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['note'], str(NOTE_CYCLE))

    def test_excel_download(self):
        """The API returns a workbook with a summary and a detail sheet."""
        url = reverse('api-stock-item-raw-materials', kwargs={'pk': self.unit.pk})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('raw_materials_Unit_U-1.xlsx', response['Content-Disposition'])

        data = b''.join(response.streaming_content)
        workbook = openpyxl.load_workbook(BytesIO(data))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Detail'])

        summary = {
            row[0]: row[3]
            for row in workbook['Summary'].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(summary['Screw'], 5)
        self.assertEqual(summary['Rope'], 1)

        detail = list(workbook['Detail'].iter_rows(min_row=2, values_only=True))
        screw = next(row for row in detail if row[6] == 'SCR-1')
        self.assertEqual(screw[5], 'Acme')
        self.assertEqual(screw[9], 'Unit #U-1 > Board (B-1)')

    def test_requires_stock_view(self):
        """Users without the stock role cannot download the breakdown."""
        self.clearRoles()
        url = reverse('api-stock-item-raw-materials', kwargs={'pk': self.unit.pk})
        self.assertEqual(self.client.get(url).status_code, 403)
